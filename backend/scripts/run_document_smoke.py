"""隔离 API/RustFS/worker smoke，可保留测试服务供浏览器验收。"""

import argparse
import asyncio
import io
import os
import re
import signal
import subprocess
import sys
import time
from pathlib import Path
from uuid import uuid4

import httpx
from docx import Document
from dotenv import dotenv_values
from redis.asyncio import Redis
from sqlalchemy.engine import make_url

root = Path(__file__).resolve().parents[2]


def run() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--serve", action="store_true")
    parser.add_argument("--real-models", action="store_true")
    args = parser.parse_args()
    values = {k: v for k, v in dotenv_values(root / ".env").items() if v is not None}
    marker = uuid4().hex[:12]
    name = "xuemian_document_" + marker + "_e2e"
    container = values["POSTGRES_CONTAINER_NAME"]
    admin = values["POSTGRES_ADMIN_USER"]
    owner = values["APP_DB_USER"]
    if not re.fullmatch(r"[A-Za-z0-9_]+", owner):
        raise RuntimeError("invalid role")
    url = make_url(values["DATABASE_URL"]).set(database=name)
    env = {
        **os.environ,
        **values,
        "DATABASE_URL": url.render_as_string(hide_password=False),
        "ENVIRONMENT": "test",
        "DOCUMENT_PROVIDER": "qwen" if args.real_models else "deterministic",
        "AUTH_ALLOWED_ORIGINS": "http://127.0.0.1:3102",
        "AUTH_REDIS_KEY_PREFIX": "xuemian:test:auth:e2e:" + marker,
        "QDRANT_URL": "http://127.0.0.1:6335",
        "QDRANT_COLLECTION": "xuemian-e2e-" + marker,
    }
    for key, suffix in [
        ("RUSTFS_QUARANTINE_BUCKET", "quarantine"),
        ("RUSTFS_DOCUMENTS_BUCKET", "documents"),
        ("RUSTFS_RECORDINGS_BUCKET", "recordings"),
        ("RUSTFS_TEMPORARY_BUCKET", "temporary"),
    ]:
        env[key] = "xuemian-e2e-" + suffix + "-" + marker
    python = str(root / "backend/.venv/bin/python")

    def sql(statement: str) -> None:
        subprocess.run(
            [
                "docker",
                "exec",
                "-i",
                container,
                "psql",
                "--username",
                admin,
                "--dbname",
                "postgres",
                "--set",
                "ON_ERROR_STOP=1",
            ],
            input=statement,
            text=True,
            check=True,
        )

    def stop_requested(*_):
        raise KeyboardInterrupt

    signal.signal(signal.SIGTERM, stop_requested)
    processes = []
    logs = []
    created = False

    def start(command: list[str], cwd: Path, process_env: dict[str, str], label: str) -> None:
        log = open("/private/tmp/xuemian-document-" + label + ".log", "w")
        logs.append(log)
        processes.append(
            subprocess.Popen(
                command,
                cwd=cwd,
                env=process_env,
                stdout=log,
                stderr=subprocess.STDOUT,
                start_new_session=True,
            )
        )

    try:
        sql(f'CREATE DATABASE "{name}" OWNER "{owner}";')
        created = True
        subprocess.run(
            [str(root / "backend/.venv/bin/alembic"), "upgrade", "head"],
            cwd=root / "backend",
            env=env,
            check=True,
        )
        subprocess.run(
            [python, "-c", "from xuemian_ai.file_management.storage_cli import run; run()"],
            cwd=root / "backend",
            env=env,
            check=True,
        )
        subprocess.run(
            [python, "-c", "from xuemian_ai.document_processing.cli import run; run()"],
            cwd=root / "backend",
            env=env,
            check=True,
        )
        start(
            [
                python,
                "-m",
                "uvicorn",
                "xuemian_ai.main:app",
                "--host",
                "127.0.0.1",
                "--port",
                "8092",
                "--no-access-log",
            ],
            root / "backend",
            env,
            "api",
        )
        start(
            [
                python,
                "-c",
                "from xuemian_ai.file_management.worker import run_worker; run_worker()",
            ],
            root / "backend",
            env,
            "file-worker",
        )
        start(
            [python, "-c", "from xuemian_ai.document_processing.worker import run; run()"],
            root / "backend",
            env,
            "worker",
        )
        with httpx.Client(trust_env=False, timeout=5) as client:
            for _ in range(60):
                try:
                    if client.get("http://127.0.0.1:8092/api/v1/health/live").status_code == 200:
                        break
                except httpx.HTTPError:
                    pass
                time.sleep(0.25)
            else:
                raise RuntimeError("isolated API unavailable")
        asyncio.run(smoke(marker, args.real_models))
        if args.serve:
            start(
                ["pnpm", "dev", "--hostname", "127.0.0.1", "--port", "3102"],
                root / "frontend",
                {**os.environ, "BACKEND_INTERNAL_URL": "http://127.0.0.1:8092"},
                "frontend",
            )
            print("ISOLATED_BROWSER_URL=http://127.0.0.1:3102", flush=True)
            while all(p.poll() is None for p in processes):
                time.sleep(1)
    finally:
        for p in processes:
            if p.poll() is None:
                os.killpg(p.pid, signal.SIGTERM)
        for p in processes:
            try:
                p.wait(timeout=10)
            except subprocess.TimeoutExpired:
                os.killpg(p.pid, signal.SIGKILL)
                p.wait()
        for log in logs:
            log.close()
        if created:
            subprocess.run(
                [python, "scripts/cleanup_file_e2e.py"], cwd=root / "backend", env=env, check=False
            )
            subprocess.run(
                [
                    python,
                    "-c",
                    (
                        "import asyncio\n"
                        "from xuemian_ai.core.config import get_settings\n"
                        "from xuemian_ai.document_processing.vector_store import VectorStore\n"
                        "async def clean():\n"
                        " s=VectorStore(get_settings())\n"
                        " try: await s.client.delete_collection(s.collection)\n"
                        " finally: await s.close()\n"
                        "asyncio.run(clean())"
                    ),
                ],
                cwd=root / "backend",
                env=env,
                check=False,
            )
            try:
                asyncio.run(clean_auth(env["REDIS_URL"], env["AUTH_REDIS_KEY_PREFIX"]))
            except Exception:
                print("isolated auth cleanup failed", flush=True)
            finally:
                sql(f'DROP DATABASE "{name}" WITH (FORCE);')
            print("isolated resources cleaned", flush=True)


async def clean_auth(url: str, prefix: str) -> None:
    if not re.fullmatch(r"xuemian:test:auth:e2e:[0-9a-f]{12}", prefix):
        raise RuntimeError("refusing non-document-smoke prefix")
    redis = Redis.from_url(url, decode_responses=True)
    try:
        async for key in redis.scan_iter(match=prefix + ":*"):
            await redis.delete(key)
    finally:
        await redis.aclose()


async def smoke(marker: str, real_models: bool = False) -> None:
    api = "http://127.0.0.1:8092/api/v1"
    async with httpx.AsyncClient(trust_env=False, timeout=10) as client:

        async def call(
            method: str,
            path: str,
            body: object | None = None,
            headers: dict[str, str] | None = None,
            status: int = 200,
        ):
            response = await client.request(method, api + path, json=body, headers=headers)
            assert response.status_code == status, (
                path,
                response.status_code,
                response.json().get("error_key"),
            )
            payload = response.json()
            assert payload.get("code", response.status_code) == response.status_code
            return payload.get("data", payload)

        async def register(username: str):
            await call(
                "POST",
                "/auth/register",
                {
                    "username": username,
                    "nickname": "Synthetic QA",
                    "password": "Synthetic-QA-2026!",
                },
                status=201,
            )
            result = await call(
                "POST", "/auth/login", {"username": username, "password": "Synthetic-QA-2026!"}
            )
            return {"Authorization": "Bearer " + result["access_token"]}

        headers = await register("doc_" + marker)
        other = await register("other_" + marker)
        await call("GET", "/users/me/ai-processing-consent", status=401)
        consent = await call("GET", "/users/me/ai-processing-consent", headers=headers)
        await call(
            "POST",
            "/users/me/ai-processing-consent",
            {"confirmed": True, "terms_version": consent["terms_version"]},
            headers,
        )
        bases = await call("GET", "/knowledge-bases", headers=headers)
        kb = bases[0]["id"]
        doc = Document()
        doc.add_heading("PostgreSQL", 1)
        doc.add_paragraph("Transactions prevent dirty reads and support atomic commits.")
        output = io.BytesIO()
        doc.save(output)
        sys.path.insert(0, str(root / "backend/tests"))
        from test_document_parsers import mixed_pdf, native_pdf, scanned_pdf

        fixtures = [
            ("notes.txt", b"PostgreSQL transaction isolation prevents dirty reads.", "text/plain"),
            (
                "notes.md",
                b"# Python\n\nPython asyncio uses an event loop to schedule coroutines.",
                "text/markdown",
            ),
            (
                "notes.docx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            ("notes.pdf", native_pdf(), "application/pdf"),
        ]
        if real_models:
            fixtures += [
                ("scan.pdf", scanned_pdf(), "application/pdf"),
                ("mixed.pdf", mixed_pdf(), "application/pdf"),
            ]
        for filename, content, mime in fixtures:
            plan = await call(
                "POST",
                f"/knowledge-bases/{kb}/file-upload-sessions",
                {"filename": filename, "size": len(content), "declared_mime": mime},
                {**headers, "Idempotency-Key": uuid4().hex},
            )
            put = await client.put(
                plan["upload_url"], content=content, headers={"Content-Type": mime}
            )
            assert put.status_code == 200
            await call(
                "POST",
                f"/file-upload-sessions/{plan['session_id']}/complete",
                {"parts": []},
                headers,
            )
            for _ in range(100):
                upload = await call(
                    "GET", f"/file-upload-sessions/{plan['session_id']}", headers=headers
                )
                if upload["status"] == "completed":
                    break
                assert upload["status"] not in {"failed", "cancelled"}, upload["status"]
                await asyncio.sleep(0.2)
            else:
                raise RuntimeError("upload validation timeout")
            file = upload["knowledge_file_id"]
            await call(
                "GET",
                f"/knowledge-bases/{kb}/files/{file}/processing-task",
                headers=other,
                status=404,
            )
            for _ in range(600):
                task = await call(
                    "GET", f"/knowledge-bases/{kb}/files/{file}/processing-task", headers=headers
                )
                if task and task["status"] == "succeeded":
                    break
                assert not task or task["status"] != "failed", (
                    task["last_error_code"] if task else None
                )
                await asyncio.sleep(0.2)
            else:
                raise RuntimeError("processing timeout")
            print(
                {
                    "fixture": filename,
                    "status": "indexed",
                    "chunks": task["active_version"]["chunk_count"],
                },
                flush=True,
            )
        result = await call(
            "POST",
            f"/knowledge-bases/{kb}/retrieval/search",
            {"query": "PostgreSQL transaction isolation", "top_n": 5},
            headers,
        )
        assert result["evidence"] and result["trace_complete"]
        print(
            {
                "http_permission_matrix": "passed",
                "formats": 4,
                "pdf_modes": 3 if real_models else 1,
                "provider": "qwen" if real_models else "deterministic",
                "retrieval": "passed",
                "trace": "complete",
            },
            flush=True,
        )


if __name__ == "__main__":
    try:
        run()
    except KeyboardInterrupt:
        pass
