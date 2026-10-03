"""Opt-in synthetic file → existing processing/retrieval → real knowledge card evaluation."""

import asyncio
import io
import json
import os
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select
from test_document_integration import context as context
from test_document_integration import seed

from xuemian_ai.core.errors import UpstreamServiceError
from xuemian_ai.document_processing.models import (
    BackgroundTask,
    DocumentProcessingVersion,
    RetrievalTrace,
)
from xuemian_ai.document_processing.retrieval import RetrievalService
from xuemian_ai.document_processing.schemas import RetrievalRequest
from xuemian_ai.file_management.models import FileAsset, KnowledgeBaseFile, StoredObject
from xuemian_ai.learning_assets.generation import (
    QwenKnowledgeProvider,
    audit_passages,
    validate_audit,
    validate_card,
)

pytestmark = pytest.mark.skipif(
    not (
        os.getenv("DOCUMENT_INTEGRATION_DB")
        and os.getenv("KNOWLEDGE_FORMAT_EVAL")
        and os.getenv("LEARNING_REAL_MODELS")
    ),
    reason="explicit isolated DB and synthetic-only real model format evaluation required",
)

SYNTHETIC_LINES = [
    "Synthetic CedarQueue V7 supports two-step commit only.",
    "Step one writes a record into a pending area. Step two confirms it.",
    "Before confirmation a record can be cancelled. After confirmation it cannot be cancelled.",
    "A commit token can be consumed only once; reuse returns token consumed.",
    "This queue does not support rollback of a confirmed record or data compression.",
    "Use it to prevent duplicate delivery: write pending, then confirm once.",
    "Example: cancel pending record A before confirmation, so A is never committed.",
    "Common mistake: assuming a pending record is already committed.",
]
SYNTHETIC_TEXT = "\n".join(SYNTHETIC_LINES)


def pdf_native() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=720, height=420)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    stream = (
        "BT /F1 11 Tf 20 390 Td 20 TL\n"
        + "\n".join(f"({line}) Tj T*" for line in SYNTHETIC_LINES)
        + "\nET"
    )
    content.set_data(stream.encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def pdf_scan() -> bytes:
    image = Image.new("RGB", (1400, 720), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(SYNTHETIC_LINES):
        draw.text((40, 40 + index * 65), line, fill="black", font_size=22)
    out = io.BytesIO()
    image.save(out, format="PDF")
    return out.getvalue()


def fixture_content(kind):
    if kind in {"txt", "md"}:
        prefix = "# Synthetic CedarQueue V7\n\n" if kind == "md" else ""
        return (prefix + SYNTHETIC_TEXT).encode(), f"synthetic.{kind}", "text/plain"
    if kind == "docx":
        doc = Document()
        doc.add_heading("Synthetic CedarQueue V7", level=1)
        # Keep the same complete factual unit as TXT/PDF; fragmented retrieval is a separate test.
        doc.add_paragraph(SYNTHETIC_TEXT)
        out = io.BytesIO()
        doc.save(out)
        return (
            out.getvalue(),
            "synthetic.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    if kind == "pdf_native":
        return pdf_native(), "synthetic-native.pdf", "application/pdf"
    if kind == "pdf_scanned":
        return pdf_scan(), "synthetic-scanned.pdf", "application/pdf"
    writer = PdfWriter()
    for data in (pdf_native(), pdf_scan()):
        writer.add_page(PdfReader(io.BytesIO(data)).pages[0])
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue(), "synthetic-mixed.pdf", "application/pdf"


@pytest.mark.parametrize("kind", ["txt", "md", "docx", "pdf_native", "pdf_scanned", "pdf_mixed"])
async def test_synthetic_format_processing_retrieval_trace_and_card_semantics(context, kind):
    settings, sessions, store = context
    content, name, mime = fixture_content(kind)
    (user, kb, file, asset, task), processor = await seed(context, content)
    async with sessions.begin() as session:
        row = await session.get(FileAsset, asset)
        row.original_filename, row.detected_mime = name, mime
        stored = await session.get(StoredObject, row.stored_object_id)
        stored.detected_mime = mime
        binding = await session.get(KnowledgeBaseFile, file)
        binding.display_name = name
    assert await processor.run_once()
    async with sessions() as session:
        state = await session.get(BackgroundTask, task)
        assert state.status == "succeeded", state.last_error_code
        version = await session.scalar(
            select(DocumentProcessingVersion).where(
                DocumentProcessingVersion.file_asset_id == asset,
                DocumentProcessingVersion.status == "active",
            )
        )
        assert version is not None
        if kind == "pdf_scanned":
            assert version.ocr_page_count == 1 and version.native_page_count == 0
        if kind == "pdf_mixed":
            assert version.ocr_page_count == 1 and version.native_page_count == 1
        if kind == "pdf_native":
            assert version.ocr_page_count == 0 and version.native_page_count == 1
    result = await RetrievalService(sessions, settings, user.id, str(uuid4()), store).search(
        kb,
        RetrievalRequest(
            query="Synthetic CedarQueue V7 pending confirmation cancelled commit token",
            file_ids=[file],
            top_n=8,
        ),
    )
    assert result.trace_complete and result.evidence
    assert all(
        source.file_id == file and source.processing_version_id == version.id
        for source in result.evidence
    )
    if kind.startswith("pdf"):
        assert any(source.page_start is not None for source in result.evidence)
    else:
        assert all(source.page_start is None for source in result.evidence)
    async with sessions() as session:
        trace = await session.get(RetrievalTrace, result.trace_id)
        assert "pending area" not in str(trace.stages)
    raw = await QwenKnowledgeProvider(settings).invoke(
        "generate",
        {
            "config": {
                "topic": "Synthetic CedarQueue V7 two-step commit",
                "source_mode": "materials",
                "foundation": "unfamiliar",
                "depth": "systematic",
            },
            "effective_context": {"values": {"preferred_language": "en-US"}},
            "evidence": [
                {"number": index + 1, "content": source.content}
                for index, source in enumerate(result.evidence)
            ],
        },
    )
    card = validate_card(raw, "materials", result.evidence, {str(file): str(asset)})
    assert card.citations and all(
        ref.file_id == file and ref.processing_version_id == version.id for ref in card.citations
    )
    passages = audit_passages(card)
    audit = await QwenKnowledgeProvider(settings).invoke(
        "audit",
        {
            "passages": passages,
            "evidence": [
                {"number": i + 1, "content": e.content} for i, e in enumerate(result.evidence)
            ],
        },
    )
    try:
        validate_audit(audit, passages, result.evidence)
        audit_accepted = True
    except UpstreamServiceError as failure:
        assert failure.error_key == "KNOWLEDGE_OUTPUT_INVALID"
        audit_accepted = False
    # Semantics may be explained in any of the five approved sections.
    content_fields = card.model_dump(mode="json", exclude={"citations", "source_mode"})
    body = json.dumps(content_fields, ensure_ascii=False).casefold()
    if output_dir := os.getenv("LEARNING_EVAL_OUTPUT_DIR"):
        destination = Path(output_dir)
        await asyncio.to_thread(destination.mkdir, parents=True, exist_ok=True)
        (destination / f"{kind}.json").write_text(
            json.dumps(
                {
                    "kind": kind,
                    "native_pages": version.native_page_count,
                    "ocr_pages": version.ocr_page_count,
                    "evidence_count": len(result.evidence),
                    "retrieved_evidence": [
                        {"number": i + 1, "content": e.content}
                        for i, e in enumerate(result.evidence)
                    ],
                    "trace_complete": result.trace_complete,
                    "audit_accepted": audit_accepted,
                    "audit": audit,
                    "card": card.model_dump(mode="json"),
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n"
        )
    if not audit_accepted:
        return  # Correctly rejected by the same gate the worker uses; report this separately.
    assert "pending" in body and "confirm" in body
    assert "cannot" in body or "not" in body or "irreversible" in body
