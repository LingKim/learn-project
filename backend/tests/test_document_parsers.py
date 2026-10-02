import io
from dataclasses import replace
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from xuemian_ai.document_processing.parsers import (
    OcrResult,
    ParserLimits,
    ProcessingError,
    chunk_nodes,
    complete_ocr,
    extract_document,
    parse_document,
)


def native_pdf() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(
        b"BT /F1 16 Tf 20 150 Td (PostgreSQL transaction isolation prevents dirty reads.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def scanned_pdf() -> bytes:
    image = Image.new("RGB", (800, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text((50, 50), "PostgreSQL transactions and isolation", fill="black", font_size=30)
    output = io.BytesIO()
    image.save(output, format="PDF")
    return output.getvalue()


def mixed_pdf() -> bytes:
    writer = PdfWriter()
    for data in (native_pdf(), scanned_pdf()):
        writer.add_page(PdfReader(io.BytesIO(data)).pages[0])
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_native_pdf_keeps_page_and_never_ocr() -> None:
    result = parse_document(native_pdf(), "pdf", ParserLimits())
    assert result.native_page_count == 1
    assert not result.ocr_pages
    assert result.blocks[0].page == 1
    assert result.blocks[0].source_kind == "native_text"


async def test_mixed_pdf_preserves_order_source_and_cleans_images() -> None:
    result = parse_document(mixed_pdf(), "pdf", ParserLimits())
    assert [p.page for p in result.ocr_pages] == [2]

    class Ocr:
        async def recognize(self, png: bytes) -> OcrResult:
            assert png.startswith(b"\x89PNG")
            return OcrResult("Scanned content contains transaction isolation and concurrency.")

    counts = []

    async def progress(done: int, total: int) -> None:
        counts.append((done, total))

    await complete_ocr(result, Ocr(), ParserLimits(), progress)
    assert [b.page for b in result.blocks] == [1, 2]
    assert [b.source_kind for b in result.blocks] == ["native_text", "ocr"]
    assert result.ocr_page_count == 1 and not result.ocr_pages
    assert result.blocks[1].confidence is None
    assert counts == [(0, 1), (1, 1)]


@pytest.mark.parametrize(
    ("text", "code"), [("", "DOCUMENT_OCR_NO_TEXT"), ("\ufffd" * 20, "DOCUMENT_OCR_LOW_CONFIDENCE")]
)
async def test_ocr_invalid_output_never_survives(text: str, code: str) -> None:
    result = parse_document(scanned_pdf(), "pdf", ParserLimits())

    class Ocr:
        async def recognize(self, png: bytes) -> OcrResult:
            return OcrResult(text)

    async def progress(done: int, total: int) -> None:
        pass

    with pytest.raises(ProcessingError, match=code):
        await complete_ocr(result, Ocr(), ParserLimits(), progress)
    assert result.ocr_pages == []


def test_pdf_render_resource_limit() -> None:
    with pytest.raises(ProcessingError, match="DOCUMENT_OCR_LIMIT_EXCEEDED"):
        parse_document(scanned_pdf(), "pdf", replace(ParserLimits(), max_pixels=10))


def test_docx_structure_keeps_heading_table_and_no_fake_pages() -> None:
    doc = Document()
    doc.add_heading("Database", level=1)
    doc.add_paragraph("Isolation levels", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Level"
    table.cell(0, 1).text = "Guarantee"
    table.cell(1, 0).text = "Serializable"
    table.cell(1, 1).text = "No anomalies"
    out = io.BytesIO()
    doc.save(out)
    result = parse_document(out.getvalue(), "docx", ParserLimits())
    assert [b.kind for b in result.blocks] == ["paragraph", "paragraph", "table"]
    assert result.blocks[-1].headings == ("Database",)
    assert "Serializable" in result.blocks[-1].text
    assert all(b.page is None for b in result.blocks)


def test_markdown_keeps_code_table_alt_and_drops_external_resources() -> None:
    source = (
        "# 数据库\n\n事务隔离。 ![示意图](https://secret.invalid/a.png)\n\n"
        "```sql\nSELECT 1;\n```\n\n| 级别 | 保证 |\n|---|---|\n| 串行化 | 无异常 |\n\n"
        "<script>alert(1)</script>"
    )
    result = parse_document(source.encode(), "md", ParserLimits())
    assert result.image_count == 1
    assert any(b.kind == "code" and "SELECT 1;" in b.text for b in result.blocks)
    assert any(b.kind == "table" for b in result.blocks)
    assert "示意图" in " ".join(b.text for b in result.blocks)
    assert "secret.invalid" not in " ".join(b.text for b in result.blocks)
    assert "alert" not in " ".join(b.text for b in result.blocks)


def test_txt_bom_paragraphs_stable_chunk_ids_and_anchors() -> None:
    result = parse_document("\ufeff事务隔离\n\n防止脏读".encode(), "txt", ParserLimits())
    version = uuid4()
    first = chunk_nodes(result, version, ParserLimits())
    second = chunk_nodes(result, version, ParserLimits())
    assert [n.node_id for n, _ in first] == [n.node_id for n, _ in second]
    assert [b.paragraph_start for _, b in first] == [1, 2]
    assert all(n.metadata["processing_version_id"] == str(version) for n, _ in first)


async def test_real_parser_subprocess_returns_and_invalid_bytes_fail_safely() -> None:
    result = await extract_document(b"PostgreSQL isolation", "notes.txt", ParserLimits())
    assert result.blocks[0].text == "PostgreSQL isolation"
    with pytest.raises(ProcessingError, match="DOCUMENT_STRUCTURE_INVALID"):
        await extract_document(b"%PDF-corrupted", "bad.pdf", ParserLimits())


def test_long_table_rows_repeat_header_without_oversize_or_tail_duplicate() -> None:
    from xuemian_ai.document_processing.parsers import ParsedDocument, StructureBlock

    document = ParsedDocument(
        blocks=[StructureBlock("Column | Meaning\n" + "x" * 300, kind="table")]
    )
    chunks = chunk_nodes(document, uuid4(), replace(ParserLimits(), chunk_size=100, overlap=10))
    assert len(chunks) == 4
    assert all(node.text.startswith("Column | Meaning\n") for node, _ in chunks)
    assert all(len(node.text) <= 100 for node, _ in chunks)
    short = ParsedDocument(blocks=[StructureBlock("x" * 95)])
    assert (
        len(chunk_nodes(short, uuid4(), replace(ParserLimits(), chunk_size=100, overlap=10))) == 1
    )


async def test_ocr_accepts_readable_mathematical_symbols_and_unicode_quotes() -> None:
    from xuemian_ai.document_processing.parsers import OcrPage, ParsedDocument

    # 与实际失败页相同的字符类别模式：正常文字 + Sm 数学符号 + Pi/Pf 引号。
    value = "状态空间与决策函数" * 10 + "策略值" + "∀∃∈∉⊂⊆⊃⊇∪∩∅∞√∑∏∫≈≠≤≥→←↔⇒⇔×÷±" + "“”"

    class Ocr:
        async def recognize(self, png: bytes) -> OcrResult:
            return OcrResult(value)

    async def progress(done: int, total: int) -> None:
        pass

    document = ParsedDocument(ocr_pages=[OcrPage(4, b"synthetic")])
    await complete_ocr(document, Ocr(), ParserLimits(), progress)
    assert document.blocks[0].text == value
    assert document.blocks[0].page == 4
    assert document.blocks[0].source_kind == "ocr"
    assert document.blocks[0].confidence is None
    assert not document.ocr_pages


@pytest.mark.parametrize(
    "value", ["∀∃→⇔“”", "\ufffd" * 20, "正文" + "\ue000" * 20, "正文" + "\x01" * 20]
)
def test_unicode_quality_gate_still_rejects_noise(value: str) -> None:
    from xuemian_ai.document_processing.parsers import valid_text

    assert not valid_text(value)
