"""受控正文提取；不读取文档指定的路径或远程资源。"""

import asyncio
import hashlib
import io
import multiprocessing
import re
import time
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass, field, replace
from multiprocessing.connection import Connection
from pathlib import PurePath
from typing import Any, Protocol
from uuid import UUID, uuid5

import pypdfium2 as pdfium  # type: ignore[import-untyped]
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from llama_index.core.schema import TextNode
from markdown_it import MarkdownIt


class ProcessingError(Exception):
    def __init__(self, code: str, *, retryable: bool = False) -> None:
        super().__init__(code)
        self.code = code
        self.retryable = retryable


@dataclass(frozen=True)
class ParserLimits:
    max_characters: int = 2_000_000
    max_pages: int = 500
    max_ocr_pages: int = 100
    dpi: int = 144
    max_pixels: int = 12_000_000
    max_image_bytes: int = 64 * 1024 * 1024
    timeout_seconds: float = 600
    chunk_size: int = 1200
    overlap: int = 120


@dataclass(frozen=True)
class StructureBlock:
    text: str = field(repr=False)
    kind: str = "paragraph"
    source_kind: str = "paragraph"
    page: int | None = None
    paragraph_start: int | None = None
    paragraph_end: int | None = None
    headings: tuple[str, ...] = ()
    confidence: float | None = None


@dataclass(frozen=True)
class OcrPage:
    page: int
    png: bytes = field(repr=False)


@dataclass
class ParsedDocument:
    blocks: list[StructureBlock] = field(default_factory=list, repr=False)
    ocr_pages: list[OcrPage] = field(default_factory=list, repr=False)
    image_count: int = 0
    native_page_count: int = 0
    ocr_page_count: int = 0


@dataclass(frozen=True)
class OcrResult:
    text: str = field(repr=False)
    confidence: float | None = None


class DocumentOcrProvider(Protocol):
    async def recognize(self, png: bytes) -> OcrResult: ...


def valid_text(value: str) -> bool:
    chars = [char for char in value if not char.isspace()]
    if not chars:
        return False
    # 技术资料的公式、箭头、弯引号和组合字符也是可读内容。
    # 至少存在文字/数字，纯符号不能充当正文；控制/私用/未分配字符仍计为噪声。
    usable = sum(
        char.isalnum() or unicodedata.category(char)[0] in {"M", "P", "S"} for char in chars
    )
    return (
        any(char.isalnum() for char in chars)
        and usable / len(chars) >= 0.8
        and "\ufffd" not in value
    )


def _paragraphs(text: str, **kwargs: Any) -> list[StructureBlock]:
    return [
        StructureBlock(part.strip(), paragraph_start=i, paragraph_end=i, **kwargs)
        for i, part in enumerate(re.split(r"\n\s*\n", text), 1)
        if part.strip()
    ]


def _pdf(content: bytes, limits: ParserLimits) -> ParsedDocument:
    result = ParsedDocument()
    doc = pdfium.PdfDocument(content)
    try:
        if len(doc) > limits.max_pages:
            raise ProcessingError("DOCUMENT_PAGE_LIMIT_EXCEEDED")
        image_bytes = 0
        seen: set[str] = set()
        for i in range(len(doc)):
            page = doc[i]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded().replace("\r\n", "\n").strip()
                finally:
                    textpage.close()
                objects = list(page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE]))
                result.image_count += len(objects)
                width, height = page.get_size()
                area = max(1, width * height)
                image_area = 0.0
                for obj in objects:
                    left, bottom, right, top = obj.get_pos()
                    image_area += max(0, right - left) * max(0, top - bottom)
                count = sum(char.isalnum() for char in text)
                native_ok = count >= 20 and valid_text(text)
                needs_ocr = bool(objects) and (
                    not text or (not native_ok and image_area / area >= 0.5)
                )
                if needs_ocr:
                    if len(result.ocr_pages) >= limits.max_ocr_pages:
                        raise ProcessingError("DOCUMENT_OCR_LIMIT_EXCEEDED")
                    scale = limits.dpi / 72
                    if width * scale * height * scale > limits.max_pixels:
                        raise ProcessingError("DOCUMENT_OCR_LIMIT_EXCEEDED")
                    bitmap = page.render(scale=scale)
                    try:
                        image = bitmap.to_pil()
                        buffer = io.BytesIO()
                        image.save(buffer, format="PNG")
                        png = buffer.getvalue()
                    finally:
                        bitmap.close()
                    image_bytes += len(png)
                    if image_bytes > limits.max_image_bytes:
                        raise ProcessingError("DOCUMENT_OCR_LIMIT_EXCEEDED")
                    result.ocr_pages.append(OcrPage(i + 1, png))
                elif text and valid_text(text):
                    digest = hashlib.sha256(text.encode()).hexdigest()
                    if len(text) < 80 or digest not in seen:
                        result.blocks.extend(
                            _paragraphs(text, page=i + 1, source_kind="native_text")
                        )
                    seen.add(digest)
                    result.native_page_count += 1
            finally:
                page.close()
    finally:
        doc.close()
    return result


def _docx(content: bytes) -> ParsedDocument:
    doc = Document(io.BytesIO(content))
    result = ParsedDocument()
    headings: list[str] = []
    paragraph = 0
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraph += 1
            value = item.text.strip()
            if not value:
                continue
            style = item.style.name if item.style else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                level = int(style[8:])
                headings = headings[: level - 1] + [value]
            result.blocks.append(
                StructureBlock(
                    value,
                    paragraph_start=paragraph,
                    paragraph_end=paragraph,
                    headings=tuple(headings),
                )
            )
        elif isinstance(item, Table):
            start = paragraph + 1
            rows = []
            for row in item.rows:
                paragraph += 1
                rows.append(" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells))
            if rows:
                result.blocks.append(
                    StructureBlock(
                        "\n".join(rows),
                        kind="table",
                        paragraph_start=start,
                        paragraph_end=paragraph,
                        headings=tuple(headings),
                    )
                )
    result.image_count = sum(1 for rel in doc.part.rels.values() if rel.reltype.endswith("/image"))
    return result


def _markdown(content: bytes) -> ParsedDocument:
    text = content.decode("utf-8-sig")
    tokens = MarkdownIt("commonmark").enable("table").parse(text)
    lines = text.splitlines()
    result = ParsedDocument()
    headings: list[str] = []
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            level = int(token.tag[1:])
            title = tokens[i + 1].content
            headings = headings[: level - 1] + [title]
        if (
            token.type
            in {
                "heading_open",
                "paragraph_open",
                "bullet_list_open",
                "ordered_list_open",
                "table_open",
                "fence",
                "code_block",
            }
            and token.map
        ):
            start, end = token.map
            kind = (
                "table"
                if token.type == "table_open"
                else "code"
                if token.type in {"fence", "code_block"}
                else "paragraph"
            )
            # 使用 AST inline 内容，删除图片 URL，只保留 alt；从不请求图片。
            section = [token]
            if token.nesting == 1:
                depth = 1
                while i + 1 < len(tokens) and depth:
                    i += 1
                    part = tokens[i]
                    section.append(part)
                    depth += part.nesting
            fragments = []
            for part in section:
                if part.type == "inline":
                    fragments.append(
                        "".join(
                            child.content if child.type not in {"softbreak", "hardbreak"} else "\n"
                            for child in (part.children or [])
                        )
                    )
                    result.image_count += sum(
                        child.type == "image" for child in (part.children or [])
                    )
            value = token.content if kind == "code" else "\n".join(fragments)
            if kind == "table":
                value = "\n".join(lines[start:end])
                value = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", value)
            if value.strip():
                result.blocks.append(
                    StructureBlock(
                        value.strip(),
                        kind=kind,
                        source_kind="markdown",
                        paragraph_start=len(result.blocks) + 1,
                        paragraph_end=len(result.blocks) + 1,
                        headings=tuple(headings),
                    )
                )
        i += 1
    return result


def parse_document(content: bytes, extension: str, limits: ParserLimits) -> ParsedDocument:
    try:
        if extension == "pdf":
            result = _pdf(content, limits)
        elif extension == "docx":
            result = _docx(content)
        elif extension == "md":
            result = _markdown(content)
        elif extension == "txt":
            if b"\x00" in content:
                raise ProcessingError("DOCUMENT_STRUCTURE_INVALID")
            result = ParsedDocument(blocks=_paragraphs(content.decode("utf-8-sig")))
        else:
            raise ProcessingError("DOCUMENT_TYPE_UNSUPPORTED")
        if sum(len(block.text) for block in result.blocks) > limits.max_characters:
            raise ProcessingError("DOCUMENT_TEXT_LIMIT_EXCEEDED")
        if not result.blocks and not result.ocr_pages:
            raise ProcessingError("DOCUMENT_NO_TEXT")
        return result
    except ProcessingError:
        raise
    except Exception:
        raise ProcessingError("DOCUMENT_STRUCTURE_INVALID") from None


def _parse_child(
    connection: Connection, content: bytes, extension: str, limits: ParserLimits
) -> None:
    try:
        connection.send(parse_document(content, extension, limits))
    except ProcessingError as exc:
        connection.send((exc.code, exc.retryable))
    except BaseException:
        connection.send(("DOCUMENT_PARSER_UNAVAILABLE", True))
    finally:
        connection.close()


async def extract_document(content: bytes, filename: str, limits: ParserLimits) -> ParsedDocument:
    """子进程超时/取消后终止并 join，不留下后台解析或临时图像。"""
    context = multiprocessing.get_context("spawn")
    reader, writer = context.Pipe(duplex=False)
    process = context.Process(
        target=_parse_child,
        args=(writer, content, PurePath(filename).suffix.lower().lstrip("."), limits),
    )
    process.start()
    writer.close()
    started = time.monotonic()
    try:
        while not reader.poll():
            if time.monotonic() - started > limits.timeout_seconds:
                raise ProcessingError("DOCUMENT_PARSER_TIMEOUT", retryable=True)
            if not process.is_alive():
                raise ProcessingError("DOCUMENT_PARSER_UNAVAILABLE", retryable=True)
            await asyncio.sleep(0.05)
        value = await asyncio.to_thread(reader.recv)
        if isinstance(value, ParsedDocument):
            return value
        raise ProcessingError(value[0], retryable=value[1])
    finally:
        reader.close()
        if process.is_alive():
            process.terminate()
        await asyncio.to_thread(process.join, 5)
        if process.is_alive():
            process.kill()
            await asyncio.to_thread(process.join)
        process.close()


async def complete_ocr(
    document: ParsedDocument,
    provider: DocumentOcrProvider,
    limits: ParserLimits,
    progress: Callable[[int, int], Any],
) -> None:
    seen = {hashlib.sha256(block.text.encode()).hexdigest() for block in document.blocks}
    total = len(document.ocr_pages)
    try:
        for index, page in enumerate(document.ocr_pages):
            await progress(index, total)
            result = await provider.recognize(page.png)
            value = result.text.strip()
            if not value:
                raise ProcessingError("DOCUMENT_OCR_NO_TEXT")
            if not valid_text(value) or (
                result.confidence is not None and result.confidence < 0.65
            ):
                raise ProcessingError("DOCUMENT_OCR_LOW_CONFIDENCE")
            digest = hashlib.sha256(value.encode()).hexdigest()
            if digest not in seen:
                document.blocks.extend(
                    _paragraphs(
                        value, page=page.page, source_kind="ocr", confidence=result.confidence
                    )
                )
            seen.add(digest)
            document.ocr_page_count += 1
            if sum(len(block.text) for block in document.blocks) > limits.max_characters:
                raise ProcessingError("DOCUMENT_TEXT_LIMIT_EXCEEDED")
            await progress(index + 1, total)
        document.blocks.sort(key=lambda block: (block.page or 0, block.paragraph_start or 0))
    finally:
        document.ocr_pages.clear()


def chunk_nodes(
    document: ParsedDocument, version_id: UUID, limits: ParserLimits
) -> list[tuple[TextNode, StructureBlock]]:
    if limits.chunk_size <= limits.overlap or limits.overlap < 0:
        raise ProcessingError("DOCUMENT_CHUNK_CONFIG_INVALID")
    result: list[tuple[TextNode, StructureBlock]] = []
    for block in document.blocks:
        values = []
        if block.kind == "table" and len(block.text) > limits.chunk_size:
            rows = block.text.splitlines()
            header, current = rows[0], rows[0]
            capacity = limits.chunk_size - len(header) - 1
            if capacity <= 0:
                raise ProcessingError("DOCUMENT_TABLE_LIMIT_EXCEEDED")
            for row in rows[1:]:
                parts = [row[start : start + capacity] for start in range(0, len(row), capacity)]
                for part in parts or [""]:
                    if len(current) + len(part) + 1 > limits.chunk_size:
                        values.append(current)
                        current = header
                    current += "\n" + part
            values.append(current)
        else:
            values = [block.text]
        for value in values:
            starts = (
                [0]
                if len(value) <= limits.chunk_size
                else range(0, len(value) - limits.overlap, limits.chunk_size - limits.overlap)
            )
            for start in starts:
                text = value[start : start + limits.chunk_size]
                node_id = str(
                    uuid5(version_id, f"{len(result)}:{hashlib.sha256(text.encode()).hexdigest()}")
                )
                result.append(
                    (
                        TextNode(
                            id_=node_id,
                            text=text,
                            metadata={"processing_version_id": str(version_id)},
                        ),
                        replace(block, text=text),
                    )
                )
                if start + limits.chunk_size >= len(value):
                    break
    if not result:
        raise ProcessingError("DOCUMENT_NO_TEXT")
    return result
